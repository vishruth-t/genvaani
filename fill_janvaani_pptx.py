from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def add_title(slide, text, left, top, width, height):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 0, 0)
    return txBox

def add_content(slide, text, left, top, width, height, font_size=24, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    
    lines = text.split('\n')
    p.text = lines[0]
    p.font.size = Pt(font_size)
    
    for line in lines[1:]:
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(font_size)
        p.alignment = align
        
    return txBox

def main():
    pptx_path = "/Users/shushil/Downloads/Sarvam/Copy of HackSprint PPT Presentation.pptx"
    out_path = "/Users/shushil/Downloads/Sarvam/JanVaani_Hackathon_Final_Presentation.pptx"
    prs = Presentation(pptx_path)
    
    # Image paths
    img_solution = "/Users/shushil/.gemini/antigravity-ide/brain/b0d6e7dd-c468-40e9-bf43-efe51bfe20c5/janvaani_solution_1791222196637.jpg"
    img_architecture = "/Users/shushil/.gemini/antigravity-ide/brain/b0d6e7dd-c468-40e9-bf43-efe51bfe20c5/janvaani_architecture_1791222208311.jpg"
    img_data_flow = "/Users/shushil/.gemini/antigravity-ide/brain/b0d6e7dd-c468-40e9-bf43-efe51bfe20c5/janvaani_data_flow_1791222220866.jpg"
    img_dashboard = "/Users/shushil/.gemini/antigravity-ide/brain/b0d6e7dd-c468-40e9-bf43-efe51bfe20c5/janvaani_dashboard_1791222233430.jpg"
    
    # SLIDE 1: Title & Team Details
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if shape.has_text_frame and "YOUR NAME:" in shape.text:
            shape.text_frame.text = "YOUR NAME: Shushil\nTEAM NAME: Alpha Voice\nTRACK NO: AI / Voice Agents\nPROJECT: JanVaani"
            for p in shape.text_frame.paragraphs:
                p.font.size = Pt(24)
                p.font.bold = True
                
    add_content(slide1, "JanVaani\nCall the Internet. Hear the Answer.", Inches(1), Inches(4.5), Inches(10), Inches(2), font_size=48)
    add_content(slide1, "A phone number anyone can call. Ask in your own Indian language.\nHear a short answer researched from multiple web sources.", Inches(1), Inches(6.5), Inches(10), Inches(2), font_size=28)
    
    # SLIDE 2: Solution
    slide2 = prs.slides[1]
    solution_text = (
        "Problem: The web requires typing, reading, and digital literacy.\n\n"
        "Solution: JanVaani\n"
        "• Phone access + Indian-language speech + live research\n"
        "• Multi-source cross-checking and source ranking\n"
        "• No app. No typing. No language barrier. Just a phone call."
    )
    add_content(slide2, solution_text, Inches(1), Inches(3), Inches(8), Inches(6), font_size=28)
    slide2.shapes.add_picture(img_solution, Inches(9.5), Inches(3), width=Inches(9))
    
    # SLIDE 3: Tech Stack & Architecture
    slide3 = prs.slides[2]
    tech_text = (
        "Core Tech Stack:\n"
        "• Telephony: Exotel (Voice streaming)\n"
        "• Gateway: FastAPI + Pipecat\n"
        "• AI Models (Sarvam AI):\n"
        "    - Saaras v4 (Real-time STT)\n"
        "    - Sarvam-105B (Orchestrator)\n"
        "    - Bulbul v3 (Streaming TTS)\n"
        "• Search Broker: Tavily + Serper\n"
        "• Database: PostgreSQL + Redis"
    )
    add_content(slide3, tech_text, Inches(1), Inches(3), Inches(8), Inches(7), font_size=28)
    slide3.shapes.add_picture(img_architecture, Inches(9.5), Inches(3), width=Inches(9))
    
    # SLIDE 4: Key Features (Title missing, add it)
    slide4 = prs.slides[3]
    add_title(slide4, "Key Features & Safety", Inches(1.12), Inches(0.91), Inches(10), Inches(1.5))
    
    features_text = (
        "• Natural Language: Accepts colloquial and code-mixed speech (e.g. Hindi, Kannada, Tamil, English).\n"
        "• Truth Grounding: Cross-checks claims across multiple sources. Never guesses if unsure.\n"
        "• Risk Router: Instantly routes Emergency or Self-Harm queries to safe paths without web search.\n"
        "• Privacy by Default: Follows DPDP Act. Explicit spoken consent required to proceed.\n"
        "• Fallback Mechanisms: If any component fails, gives a useful fallback, never silent hallucination."
    )
    add_content(slide4, features_text, Inches(1), Inches(3), Inches(18), Inches(6), font_size=32)
    
    # SLIDE 5: Data Flow Diagram
    slide5 = prs.slides[4]
    flow_text = (
        "1. Caller dials Exotel number -> Streams 8kHz audio to Voice Gateway.\n"
        "2. Sarvam Saaras transcribes and detects language.\n"
        "3. Orchestrator checks intent & risk -> Search Broker queries Tavily/Serper.\n"
        "4. Evidence Extractor filters, deduplicates, and corroborates claims.\n"
        "5. Sarvam-105B synthesizes a grounded answer.\n"
        "6. Sarvam Bulbul speaks the answer back in the caller's language."
    )
    add_content(slide5, flow_text, Inches(1), Inches(2.5), Inches(8.5), Inches(7), font_size=26)
    slide5.shapes.add_picture(img_data_flow, Inches(10), Inches(2.5), width=Inches(8.5))
    
    # SLIDE 6: Screenshot of Your Project
    slide6 = prs.slides[5]
    screenshot_text = (
        "JanVaani Live Dashboard\n\n"
        "Monitoring real-time call transcripts, source verifications, latency metrics, and confidence scores."
    )
    add_content(slide6, screenshot_text, Inches(1), Inches(3), Inches(8), Inches(6), font_size=28)
    slide6.shapes.add_picture(img_dashboard, Inches(9.5), Inches(3), width=Inches(9))
    
    # SLIDE 7: Others
    slide7 = prs.slides[6]
    for shape in slide7.shapes:
        if shape.has_text_frame and "Others" in shape.text:
            shape.text_frame.text = "Impact & Roadmap"
            
    impact_text = (
        "Evaluation & Metrics:\n"
        "- Target 85%+ answer correctness & 95%+ groundedness.\n"
        "- Safe handling of all medical/legal/emergency prompts.\n\n"
        "Future Roadmap:\n"
        "- Multi-turn Context: Remembering previous questions in a session.\n"
        "- Follow-up Actions: Send the source links to the caller via SMS or WhatsApp.\n"
        "- Real-time Vector Search: For private enterprise data."
    )
    add_content(slide7, impact_text, Inches(1), Inches(3), Inches(18), Inches(6), font_size=32)
    
    prs.save(out_path)
    print(f"Successfully saved to {out_path}")

if __name__ == '__main__':
    main()
