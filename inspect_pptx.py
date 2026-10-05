from pptx import Presentation
from pptx.util import Inches, Pt

pptx_path = "/Users/shushil/Downloads/Sarvam/Copy of HackSprint PPT Presentation.pptx"
prs = Presentation(pptx_path)

for i, slide in enumerate(prs.slides):
    print(f"--- Slide {i+1} ---")
    for j, shape in enumerate(slide.shapes):
        print(f"  Shape {j}: {shape.name}")
        if shape.has_text_frame:
            print(f"    Text: {shape.text.replace(chr(10), ' ')}")
            print(f"    Left: {shape.left}, Top: {shape.top}, Width: {shape.width}, Height: {shape.height}")
