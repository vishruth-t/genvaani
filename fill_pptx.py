from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def add_title(slide, text, left, top, width, height):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 0, 0)
    return txBox

def add_content(slide, text, left, top, width, height, font_size=18, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    
    lines = text.split('\n')
    p.text = lines[0]
    p.font.size = Pt(font_size)
    
    for line in lines[1:]:
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(font_size)
        p.alignment = align
        
    return txBox

def main():
    pptx_path = "/Users/shushil/Downloads/Sarvam/Copy of HackSprint PPT Presentation.pptx"
    out_path = "/Users/shushil/Downloads/Sarvam/Sarvam_Voice_Agent_Hackathon_Final_Presentation.pptx"
    prs = Presentation(pptx_path)
    
    # EMUs to Inches converter helper
    def to_inches(emu):
        return emu / 914400.0
    
    # SLIDE 1: Title & Team Details
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if shape.has_text_frame and "YOUR NAME:" in shape.text:
            shape.text_frame.text = "YOUR NAME: Shushil\nTEAM NAME: Alpha Voice\nTRACK NO: AI / Voice Agents\nPROJECT: Sarvam Voice Agent"
            for p in shape.text_frame.paragraphs:
                p.font.size = Pt(20)
                p.font.bold = True
                
    # We should add the Project Name to Slide 1 prominently.
    add_content(slide1, "Sarvam Voice Agent\nReal-time, multilingual AI phone agent", Inches(0.5), Inches(3.5), Inches(7), Inches(2), font_size=36)
    
    # SLIDE 2: Solution
    slide2 = prs.slides[1]
    solution_text = (
        "Problem Statement:\n"
        "Customer support centers face high latency, language barriers, and lack of context-awareness with traditional voice bots. Voice conversion overhead (e.g., PCM to mu-law) adds significant delay.\n\n"
        "Our Solution:\n"
        "- A production-quality, voice-based AI phone agent built using Twilio Media Streams.\n"
        "- Zero Audio Conversion: Uses raw 8kHz mu-law audio end-to-end for ultra-low latency.\n"
        "- Multilingual Support: Auto-detects English, Hindi, and Odia and switches language seamlessly.\n"
        "- Context-Aware & Actionable: Integrated with RAG for knowledge and tools for actions like booking appointments."
    )
    add_content(slide2, solution_text, Inches(1), Inches(1.5), Inches(8), Inches(4), font_size=20)
    
    # SLIDE 3: Tech Stack & Architecture
    slide3 = prs.slides[2]
    tech_text = (
        "Tech Stack:\n"
        "• Language: Python 3.9+\n"
        "• Framework: FastAPI (Async WebSockets)\n"
        "• Telephony: Twilio Media Streams\n"
        "• AI Models (Sarvam AI Suite):\n"
        "    - STT: saaras:v3-realtime / saaras:v4 (Real-time Speech-to-Text)\n"
        "    - LLM: sarvam-105b-conversations (Chat Completions)\n"
        "    - TTS: bulbul:v3 (Streaming Text-to-Speech)\n\n"
        "Architecture Flow:\n"
        "User Phone -> Twilio Stream -> FastAPI Server -> Sarvam STT -> LLM Engine + RAG -> Sarvam TTS -> Twilio Stream -> User"
    )
    add_content(slide3, tech_text, Inches(1), Inches(1.5), Inches(8), Inches(4.5), font_size=20)
    
    # SLIDE 4: Key Features
    slide4 = prs.slides[3]
    # Add title since it's missing
    # Title from slide 3 was: Left: 1028700 (1.12"), Top: 838200 (0.91")
    add_title(slide4, "Key Features", Inches(1.12), Inches(0.91), Inches(10), Inches(1))
    
    features_text = (
        "• Ultra-Low Latency: Direct mu-law streaming without intermediate transcoding overhead.\n"
        "• Intelligent Barge-in: Gracefully handles user interruptions by clearing audio buffers and instantly resetting conversation state.\n"
        "• Smart Silence Handling: Detects user silence and automatically prompts them to continue.\n"
        "• Dynamic Tool Calling: The agent can invoke backend APIs mid-conversation (e.g., lookup order status, transfer to human).\n"
        "• Localized RAG Integration: Queries local knowledge bases to ground responses in facts."
    )
    add_content(slide4, features_text, Inches(1), Inches(1.5), Inches(8), Inches(4), font_size=20)
    
    # SLIDE 5: Data Flow Diagram
    slide5 = prs.slides[4]
    flow_text = (
        "1. Ingestion: Twilio forwards WebSocket media payload (8kHz mu-law) to FastAPI.\n"
        "2. Transcription: Audio is queued and streamed to Sarvam STT. Auto-language detection identifies if user speaks Hindi, English, or Odia.\n"
        "3. Orchestration & RAG: The Agent intercepts 'transcript.final', searches the RAG KnowledgeBase, and appends context.\n"
        "4. Reasoning: The LLM streams completions. If a tool is called, the Agent executes it synchronously.\n"
        "5. Speech Generation: Text is chunked based on sentence boundaries and streamed to Sarvam TTS.\n"
        "6. Playback: TTS returns mu-law audio chunks which are relayed directly back to Twilio."
    )
    add_content(slide5, flow_text, Inches(1), Inches(1.5), Inches(8), Inches(4.5), font_size=18)
    
    # SLIDE 6: Screenshot of Your Project
    slide6 = prs.slides[5]
    screenshot_text = (
        "Voice Agent Interface (Console & TwiML)\n\n"
        "Since this is a headless voice application, the primary interfaces are telephony and backend logs.\n\n"
        "TwiML Configuration for Twilio:\n"
        "<Response>\n"
        "    <Connect>\n"
        "        <Stream url=\"wss://<ngrok-id>.ngrok-free.app/media\" />\n"
        "    </Connect>\n"
        "</Response>\n\n"
        "Backend Execution Log Example:\n"
        "INFO: Connecting to Sarvam STT...\n"
        "INFO: Received start event stream_sid=xyz\n"
        "INFO: User said text='mera order kahan hai' lang='hi-IN'\n"
        "INFO: Tool called: lookup_order_status\n"
        "INFO: Agent said text='आपका ऑर्डर कल डिलीवर हो जाएगा।'\n"
    )
    add_content(slide6, screenshot_text, Inches(1), Inches(1.5), Inches(8), Inches(4.5), font_size=16)
    
    # SLIDE 7: Others
    slide7 = prs.slides[6]
    for shape in slide7.shapes:
        if shape.has_text_frame and "Others" in shape.text:
            shape.text_frame.text = "Impact & Future Scope"
            
    impact_text = (
        "Expected Impact:\n"
        "- Reduced Call Center Costs: Automates routine tasks like order lookups and booking.\n"
        "- Improved CSAT: Natural, multi-lingual, and low-latency interaction improves customer satisfaction.\n"
        "- High Accessibility: Regional language support (Hindi, Odia) broadens user base in rural demographics.\n\n"
        "Future Scope:\n"
        "- Vector Database Integration: Replace simple mock RAG with Pinecone/Milvus for large-scale enterprise knowledge.\n"
        "- Sentiment Analysis: Detect frustrated users and auto-transfer to human agents.\n"
        "- More Tools: CRM integration (Salesforce/Hubspot) for real-world appointment scheduling."
    )
    add_content(slide7, impact_text, Inches(1), Inches(1.5), Inches(8), Inches(4.5), font_size=18)
    
    prs.save(out_path)
    print(f"Successfully saved to {out_path}")

if __name__ == '__main__':
    main()
